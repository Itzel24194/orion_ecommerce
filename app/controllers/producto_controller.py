# ================================================================
# app/controllers/producto_controller.py - CONTROLADOR COMPLETO
# ================================================================

import os
import uuid
import csv
import urllib.request
import traceback
from io import StringIO, BytesIO
from flask import render_template, request, redirect, url_for, current_app, session, jsonify, make_response, flash, Response
from datetime import datetime
from app.models.marcas_model import Marca
from app.models.resenas_model import Resena
from app.models.categorias_model import Categoria
from app.models.productos_model import Producto as ProductoModel
from app.models.temporada_model import Temporada
from bson import ObjectId
from werkzeug.utils import secure_filename

# ✅ importación de Excel
try:
    import openpyxl
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    OPENPYXL_OK = True
except ImportError:
    OPENPYXL_OK = False
    print("⚠️ openpyxl no instalado. Ejecuta: python -m pip install openpyxl")

# ⭐ IMPORT MARKETPLACE (opcional)
try:
    from app.models.marketplace_model import Vendedor
    MARKETPLACE_OK = True
except ImportError:
    Vendedor = None
    MARKETPLACE_OK = False
    print("ℹ️ Módulo Marketplace no disponible (opcional)")


# --- UTILIDADES ---

EXTENSIONES_IMAGEN = [
    '.jpg', '.jpeg', '.jpe', '.jfif', '.pjpeg', '.pjp',
    '.png', '.apng', '.gif', '.webp', '.bmp', '.dib',
    '.tif', '.tiff', '.avif', '.heic', '.heif', '.svg',
]

EXTENSIONES_VIDEO = [
    '.mp4', '.m4v', '.webm', '.ogv', '.ogg',
    '.mov', '.qt', '.avi', '.mkv', '.flv',
    '.wmv', '.3gp', '.3g2', '.mpeg', '.mpg', '.m2v',
]

CONTENT_TYPE_A_EXT_IMG = {
    'image/jpeg': '.jpg', 'image/jpg': '.jpg', 'image/pjpeg': '.jpg',
    'image/png': '.png', 'image/gif': '.gif', 'image/webp': '.webp',
    'image/bmp': '.bmp', 'image/tiff': '.tiff', 'image/avif': '.avif',
    'image/heic': '.heic', 'image/heif': '.heif', 'image/svg+xml': '.svg',
}

CONTENT_TYPE_A_EXT_VID = {
    'video/mp4': '.mp4', 'video/x-m4v': '.m4v', 'video/webm': '.webm',
    'video/ogg': '.ogv', 'video/quicktime': '.mov', 'video/x-msvideo': '.avi',
    'video/x-matroska': '.mkv', 'video/x-flv': '.flv', 'video/x-ms-wmv': '.wmv',
    'video/3gpp': '.3gp', 'video/mpeg': '.mpeg',
}

MAX_FOTOS = 6
MAX_VIDEOS = 3
MAX_SIZE_VIDEO = 100 * 1024 * 1024

BADGE_VENDEDOR_MAX_LEN = 20


def _get_upload_folder():
    folder = os.path.join(current_app.root_path, 'static', 'uploads', 'productos')
    os.makedirs(folder, exist_ok=True)
    return folder


def _save_photos(files, nombre_producto, fotos_existentes=None):
    upload_path = _get_upload_folder()
    prefijo = "".join([c if c.isalnum() else "_" for c in (nombre_producto or 'producto').lower()])[:30]
    fotos_guardadas = list(fotos_existentes) if fotos_existentes else []

    archivos_lista = list(files) if files else []
    print(f"   📎 _save_photos: recibidos {len(archivos_lista)} archivo(s)")

    for idx, file in enumerate(archivos_lista):
        fname = getattr(file, 'filename', None)
        ctype = getattr(file, 'content_type', None)

        if file is None or not fname or not str(fname).strip():
            continue
        if len(fotos_guardadas) >= MAX_FOTOS:
            break

        ext_original = os.path.splitext(fname)[1].lower()
        if ext_original in EXTENSIONES_IMAGEN:
            ext = ext_original
        else:
            ext = CONTENT_TYPE_A_EXT_IMG.get((ctype or '').lower(), '.jpg')

        nombre_archivo = f"{prefijo}_img_{len(fotos_guardadas)}_{os.urandom(2).hex()}{ext}"
        ruta_destino = os.path.join(upload_path, nombre_archivo)

        try:
            try:
                file.seek(0)
            except Exception:
                pass
            file.save(ruta_destino)
            if os.path.exists(ruta_destino) and os.path.getsize(ruta_destino) > 0:
                fotos_guardadas.append(nombre_archivo)
                print(f"         ✅ Guardado: {nombre_archivo}")
        except Exception as e:
            print(f"         ❌ Error: {e}")
            traceback.print_exc()

    return fotos_guardadas


def _save_videos(files, nombre_producto, videos_existentes=None):
    upload_path = _get_upload_folder()
    prefijo = "".join([c if c.isalnum() else "_" for c in (nombre_producto or 'producto').lower()])[:30]
    videos_guardados = list(videos_existentes) if videos_existentes else []

    archivos_lista = list(files) if files else []
    print(f"   🎬 _save_videos: recibidos {len(archivos_lista)} archivo(s)")

    for idx, file in enumerate(archivos_lista):
        fname = getattr(file, 'filename', None)
        ctype = getattr(file, 'content_type', None)

        if file is None or not fname or not str(fname).strip():
            continue
        if len(videos_guardados) >= MAX_VIDEOS:
            break

        ext_original = os.path.splitext(fname)[1].lower()
        if ext_original in EXTENSIONES_VIDEO:
            ext = ext_original
        else:
            ext = CONTENT_TYPE_A_EXT_VID.get((ctype or '').lower(), '.mp4')

        nombre_archivo = f"{prefijo}_vid_{len(videos_guardados)}_{os.urandom(2).hex()}{ext}"
        ruta_destino = os.path.join(upload_path, nombre_archivo)

        try:
            try:
                file.seek(0)
            except Exception:
                pass
            file.save(ruta_destino)
            if os.path.exists(ruta_destino):
                size = os.path.getsize(ruta_destino)
                if size > MAX_SIZE_VIDEO:
                    os.remove(ruta_destino)
                    continue
                if size > 0:
                    videos_guardados.append(nombre_archivo)
                    print(f"         ✅ Guardado: {nombre_archivo}")
        except Exception as e:
            print(f"         ❌ Error: {e}")

    return videos_guardados


def _extraer_variantes(form):
    skus = form.getlist('sku[]')
    colores = form.getlist('color[]')
    tamanos = form.getlist('tamano[]')
    precios = form.getlist('precio[]')
    stocks = form.getlist('stock[]')
    variantes = []
    for i in range(len(skus)):
        if skus[i].strip():
            variantes.append({
                "sku": skus[i].strip(),
                "color": colores[i].strip() if i < len(colores) else "",
                "tamano": tamanos[i].strip() if i < len(tamanos) else "",
                "precio": float(precios[i]) if i < len(precios) and precios[i] not in (None, "") else 0.0,
                "stock": int(stocks[i]) if i < len(stocks) and stocks[i] not in (None, "") else 0
            })
    return variantes


# ====================================================================
# ⭐ MARKETPLACE — HELPERS
# ====================================================================
def _get_vendedor_actual():
    if not MARKETPLACE_OK or Vendedor is None:
        return None
    if not session.get('user_id'):
        return None
    try:
        db = current_app.db
        v = Vendedor.obtener_por_user(db, session.get('user_id'))
        if v and v.get('estado') == 'aprobado' and v.get('activo', True):
            return v
    except Exception as e:
        print(f"⚠️ [_get_vendedor_actual] Error: {e}")
    return None


def _es_admin_actual():
    return session.get('rol') == 'admin'


def _truncar_nombre_vendedor(nombre, max_len=BADGE_VENDEDOR_MAX_LEN):
    if not nombre:
        return ""
    nombre = str(nombre).strip()
    if len(nombre) <= max_len:
        return nombre
    return nombre[:max_len].rstrip() + "…"


def _nombre_vendedor_display(producto):
    """Replica la cascada del Jinja."""
    if not producto:
        return ""

    v = producto.get('vendedor_nombre')
    if v:
        return str(v).strip()

    v = producto.get('nombre_tienda')
    if v:
        return str(v).strip()

    v = producto.get('tienda_nombre')
    if v:
        return str(v).strip()

    vendedor_raw = producto.get('vendedor')

    if isinstance(vendedor_raw, str) and vendedor_raw.strip():
        return vendedor_raw.strip()

    if isinstance(vendedor_raw, dict):
        for k in ('nombre_tienda', 'tienda_nombre', 'nombre', 'name'):
            if vendedor_raw.get(k):
                return str(vendedor_raw[k]).strip()

    vinfo = producto.get('vendedor_info') or {}
    if isinstance(vinfo, dict) and vinfo.get('nombre'):
        return str(vinfo['nombre']).strip()

    return ""


def _aplicar_campos_planos_vendedor(producto, nombre, vendedor_id_str, slug):
    """Aplica los campos planos que el template Jinja espera."""
    nombre = (nombre or '').strip()

    producto['vendedor_nombre'] = nombre
    producto['nombre_tienda']   = nombre
    producto['tienda_nombre']   = nombre
    producto['vendedor']        = nombre

    producto['vendedor_id_str']       = vendedor_id_str
    producto['vendedor_slug']         = slug or ''
    producto['vendedor_nombre_badge'] = _truncar_nombre_vendedor(
        nombre, BADGE_VENDEDOR_MAX_LEN
    )

    producto['data_vendidos'] = slug or vendedor_id_str or nombre


def _enriquecer_producto_con_vendedor(producto):
    """Agrega info del vendedor + campos planos para el badge."""
    if not producto:
        return producto

    vendedor_id = producto.get('vendedor_id')

    if not vendedor_id:
        producto['vendedor_info']  = None
        producto['es_marketplace'] = False
        producto.setdefault('vendedor_nombre', '')
        return producto

    producto['es_marketplace'] = True

    nombre_fallback = (
        producto.get('vendedor_nombre')
        or producto.get('nombre_tienda')
        or producto.get('tienda_nombre')
        or ''
    )

    if not MARKETPLACE_OK or Vendedor is None:
        if nombre_fallback:
            producto['vendedor_info'] = {
                'id':     str(vendedor_id),
                'nombre': nombre_fallback,
                'slug':   None,
                'logo':   '',
            }
        else:
            producto['vendedor_info'] = None
        _aplicar_campos_planos_vendedor(producto, nombre_fallback, str(vendedor_id), None)
        return producto

    try:
        db = current_app.db
        v = Vendedor.obtener_por_id(db, vendedor_id)

        if v and v.get('estado') == 'aprobado':
            nombre = v.get('nombre_tienda') or nombre_fallback or 'Vendedor'
            producto['vendedor_info'] = {
                'id':     str(v['_id']),
                'nombre': nombre,
                'slug':   v.get('slug'),
                'logo':   v.get('logo', ''),
            }
            _aplicar_campos_planos_vendedor(
                producto, nombre,
                str(v['_id']),
                v.get('slug')
            )
        else:
            producto['vendedor_info'] = None
            if nombre_fallback:
                _aplicar_campos_planos_vendedor(
                    producto, nombre_fallback, str(vendedor_id), None
                )
    except Exception as e:
        print(f"⚠️ [_enriquecer_producto_con_vendedor] Error: {e}")
        if nombre_fallback:
            producto['vendedor_info'] = {
                'id':     str(vendedor_id),
                'nombre': nombre_fallback,
                'slug':   None,
                'logo':   '',
            }
            _aplicar_campos_planos_vendedor(
                producto, nombre_fallback, str(vendedor_id), None
            )
        else:
            producto['vendedor_info'] = None

    return producto


def _enriquecer_lista_con_vendedor(productos):
    for p in productos or []:
        _enriquecer_producto_con_vendedor(p)
    return productos


def _actualizar_metricas_vendedor(vendedor_id, subtotal, comision_pct=None):
    if not MARKETPLACE_OK or Vendedor is None or not vendedor_id:
        return False
    try:
        db = current_app.db
        v = Vendedor.obtener_por_id(db, vendedor_id)
        if not v:
            return False

        pct = comision_pct if comision_pct is not None else float(v.get('comision_pct', 15))
        comision = float(subtotal) * (pct / 100.0)

        db[Vendedor.COLLECTION].update_one(
            {'_id': v['_id']},
            {
                '$inc': {
                    'metricas.ventas':     1,
                    'metricas.ingresos':   float(subtotal),
                    'metricas.comisiones': comision,
                },
                '$set': {'updated_at': datetime.utcnow()},
            }
        )
        return True
    except Exception as e:
        print(f"⚠️ [_actualizar_metricas_vendedor] Error: {e}")
        return False


# ====================================================================
# Jerarquía de categorías
# ====================================================================
def _calcular_jerarquia_categorias(categorias_crudas):
    todas = list(categorias_crudas)
    mapa = {str(c['_id']): c for c in todas}

    for c in todas:
        padre_id = c.get('padre_id')
        if not padre_id or str(padre_id) in ['None', 'null', '']:
            c['nivel'] = 1
            c['padre_nombre'] = "Raíz"
            c['padre_id_str'] = None
        else:
            padre_id_str = str(padre_id)
            padre = mapa.get(padre_id_str)
            if padre:
                abuelo_id = padre.get('padre_id')
                if abuelo_id and str(abuelo_id) not in ['None', 'null', '']:
                    c['nivel'] = 3
                else:
                    c['nivel'] = 2
                c['padre_nombre'] = padre.get('nombre', 'Desconocido')
                c['padre_id_str'] = padre_id_str
            else:
                c['nivel'] = 1
                c['padre_nombre'] = "Raíz (Corregido)"
                c['padre_id_str'] = None

    return todas


# ====================================================================
# Favoritos
# ====================================================================
def _get_favoritos_ids():
    if not session.get('user_id'):
        return set()
    try:
        db = current_app.db
        usuario_id = ObjectId(session['user_id'])
        favs = db.favoritos.find({'usuario_id': usuario_id}, {'producto_id': 1})
        return {str(f['producto_id']) for f in favs}
    except Exception:
        return set()


# ====================================================================
# ⭐ HELPERS ROBUSTOS DE CARGA DE PRODUCTOS
# ====================================================================

def _obtener_activos_incluyendo_marketplace():
    """
    Trae TODOS los productos visibles en el catálogo (incluye marketplace).
    Query simple + filtro de estados en Python.
    """
    db = current_app.db

    productos = list(db.productos.find({'activo': {'$ne': False}}))

    ESTADOS_EXCLUIDOS = {'inactivo', 'borrador', 'deshabilitado', 'eliminado', 'rechazado'}
    productos = [
        p for p in productos
        if (p.get('estado') or 'activo').lower() not in ESTADOS_EXCLUIDOS
    ]

    print(f"🔎 [_obtener_activos_incluyendo_marketplace] {len(productos)} productos encontrados")
    con_v = sum(1 for p in productos if p.get('vendedor_id'))
    print(f"   ├─ con vendedor_id: {con_v}")
    print(f"   └─ sin vendedor_id: {len(productos) - con_v}")

    return productos


def _obtener_activos_por_categoria(cat_id):
    """
    Carga TODOS los productos activos y filtra por categoría + descendientes
    en Python (comparación de strings). Incluye diagnóstico detallado si
    no encuentra nada.
    """
    productos = _obtener_activos_incluyendo_marketplace()

    categorias_crudas = Categoria.obtener_todas()

    # Recolectar la categoría + TODOS sus descendientes
    ids_categoria = {str(cat_id)}
    cambio = True
    while cambio:
        cambio = False
        for c in categorias_crudas:
            padre = c.get('padre_id')
            if padre and str(padre) in ids_categoria and str(c['_id']) not in ids_categoria:
                ids_categoria.add(str(c['_id']))
                cambio = True

    print(f"🔎 [_obtener_activos_por_categoria] buscando en {len(ids_categoria)} categorías: "
          f"{sorted(ids_categoria)[:5]}...")
    print(f"   total productos a filtrar: {len(productos)}")

    def _match_categoria(p):
        # Probar todos los campos posibles
        for campo in ('categoria_id', 'categoria', 'categoria_ids'):
            val = p.get(campo)
            if val is None:
                continue
            if isinstance(val, list):
                for v in val:
                    if str(v) in ids_categoria:
                        return True
            else:
                if str(val) in ids_categoria:
                    return True
        return False

    filtrados = [p for p in productos if _match_categoria(p)]

    print(f"🔎 [_obtener_activos_por_categoria] cat={cat_id} "
          f"→ {len(filtrados)} de {len(productos)} productos")

    # ⭐ Diagnóstico: si no hay ninguno, mostrar qué categorías SÍ tienen
    if not filtrados and productos:
        cats_de_productos = {}
        for p in productos:
            cid = p.get('categoria_id') or p.get('categoria')
            if cid:
                cid_str = str(cid)
                cats_de_productos[cid_str] = cats_de_productos.get(cid_str, 0) + 1
        print(f"   ⚠️ NINGÚN producto matchea. Categorías reales de los productos:")
        for cid, n in list(cats_de_productos.items())[:10]:
            print(f"      → {cid}: {n} producto(s)")

        print(f"   ⚠️ Productos de vendedor y sus categorías:")
        for p in productos:
            if p.get('vendedor_id'):
                cid = p.get('categoria_id') or p.get('categoria')
                print(f"      → '{p.get('nombre')}' cat={cid}")

    return filtrados


# --- CONTROLADORES ---

def listar_productos():
    productos = ProductoModel.obtener_todos()
    categorias_crudas = Categoria.obtener_todas()
    categorias = _calcular_jerarquia_categorias(categorias_crudas)
    marcas = Marca.obtener_todas()
    temporadas = Temporada.obtener_todas()
    temporadas_dict = {str(t['_id']): t for t in temporadas}

    vendedor_actual = _get_vendedor_actual()
    es_admin = _es_admin_actual()
    if vendedor_actual and not es_admin:
        vid_str = str(vendedor_actual['_id'])
        productos = [p for p in productos if str(p.get('vendedor_id') or '') == vid_str]

    for p in productos:
        p['marca_nombre'] = ''
        for marca in marcas:
            if str(marca.get('_id')) == str(p.get('marca_id')):
                p['marca_nombre'] = marca.get('nombre_comercial', '')
                break
        _enriquecer_producto_con_vendedor(p)

    return render_template('admin/productos.html',
                           productos=productos,
                           categorias=categorias,
                           marcas=marcas,
                           temporadas=temporadas,
                           temporadas_dict=temporadas_dict,
                           producto_seleccionado=None,
                           vendedor_actual=vendedor_actual)


def ver_producto(id):
    producto = ProductoModel.obtener_por_id(id)
    if not producto:
        flash('Producto no encontrado', 'danger')
        return redirect(url_for('web.lista_productos'))

    productos = ProductoModel.obtener_todos()
    categorias_crudas = Categoria.obtener_todas()
    categorias = _calcular_jerarquia_categorias(categorias_crudas)
    marcas = Marca.obtener_todas()
    temporadas = Temporada.obtener_todas()
    temporadas_dict = {str(t['_id']): t for t in temporadas}

    for p in productos:
        p['marca_nombre'] = ''
        for marca in marcas:
            if str(marca.get('_id')) == str(p.get('marca_id')):
                p['marca_nombre'] = marca.get('nombre_comercial', '')
                break

    for marca in marcas:
        if str(marca.get('_id')) == str(producto.get('marca_id')):
            producto['marca_nombre'] = marca.get('nombre_comercial', '')
            break

    _enriquecer_producto_con_vendedor(producto)

    vendedor_actual = _get_vendedor_actual()

    return render_template('admin/productos.html',
                           productos=productos,
                           categorias=categorias,
                           marcas=marcas,
                           temporadas=temporadas,
                           temporadas_dict=temporadas_dict,
                           producto_seleccionado=producto,
                           vendedor_actual=vendedor_actual)


def agregar():
    if request.method == 'POST':
        nombre = request.form.get('nombre')
        if not nombre:
            flash('El nombre del producto es requerido', 'danger')
            return redirect(url_for('web.lista_productos'))

        variantes = _extraer_variantes(request.form)
        archivos_img = request.files.getlist('fotos[]')
        fotos = _save_photos(archivos_img, nombre)
        archivos_vid = request.files.getlist('videos[]')
        videos = _save_videos(archivos_vid, nombre)

        temporada_id = request.form.get('temporada_id', '').strip() or None

        data = {
            "nombre": nombre.strip(),
            "descripcion": request.form.get('descripcion', '').strip(),
            "categoria_id": request.form.get('categoria_id', '').strip() or None,
            "marca_id": request.form.get('marca_id', '').strip() or None,
            "temporada_id": temporada_id,
            "estado": request.form.get('estado', 'activo').strip(),
            "activo": True,
            "variables": variantes,
            "fotos": fotos,
            "videos": videos,
            "created_at": datetime.utcnow(),
            "updated_at": datetime.utcnow()
        }

        vendedor_actual = _get_vendedor_actual()
        if vendedor_actual:
            data['vendedor_id']     = str(vendedor_actual['_id'])
            data['vendedor_nombre'] = vendedor_actual.get('nombre_tienda', '')
            data['nombre_tienda']   = vendedor_actual.get('nombre_tienda', '')
            data['tienda_nombre']   = vendedor_actual.get('nombre_tienda', '')

            try:
                db = current_app.db
                db[Vendedor.COLLECTION].update_one(
                    {'_id': vendedor_actual['_id']},
                    {'$inc': {'metricas.productos': 1}}
                )
            except Exception as e:
                print(f"   ⚠️ No se pudo actualizar métrica: {e}")

        ProductoModel.crear(data)
        flash(f'Producto agregado: {len(fotos)} foto(s), {len(videos)} video(s)', 'success')
        return redirect(url_for('web.lista_productos'))

    return redirect(url_for('web.lista_productos'))


def editar(id):
    producto = ProductoModel.obtener_por_id(id)
    if not producto:
        flash('Producto no encontrado', 'danger')
        return redirect(url_for('web.lista_productos'))

    if request.method == 'POST':
        nombre = request.form.get('nombre')
        if not nombre:
            flash('El nombre del producto es requerido', 'danger')
            return redirect(url_for('web.ver_producto', id=id))

        variantes = _extraer_variantes(request.form)

        fotos_existentes = request.form.getlist('fotos_existentes[]')
        archivos_img = request.files.getlist('fotos[]')
        fotos = _save_photos(archivos_img, nombre, fotos_existentes)

        videos_existentes = request.form.getlist('videos_existentes[]')
        archivos_vid = request.files.getlist('videos[]')
        videos = _save_videos(archivos_vid, nombre, videos_existentes)

        temporada_id = request.form.get('temporada_id', '').strip() or None

        data = {
            "nombre": nombre.strip(),
            "descripcion": request.form.get('descripcion', '').strip(),
            "categoria_id": request.form.get('categoria_id', '').strip() or None,
            "marca_id": request.form.get('marca_id', '').strip() or None,
            "temporada_id": temporada_id,
            "estado": request.form.get('estado', 'activo').strip(),
            "variables": variantes,
            "fotos": fotos,
            "videos": videos,
            "updated_at": datetime.utcnow()
        }

        if producto.get('vendedor_id'):
            data['vendedor_id']     = producto['vendedor_id']
            data['vendedor_nombre'] = producto.get('vendedor_nombre', '')
            data['nombre_tienda']   = producto.get('nombre_tienda', producto.get('vendedor_nombre', ''))
            data['tienda_nombre']   = producto.get('tienda_nombre', producto.get('vendedor_nombre', ''))

        ProductoModel.actualizar(id, data)
        flash(f'Producto actualizado: {len(fotos)} foto(s), {len(videos)} video(s)', 'success')
        return redirect(url_for('web.lista_productos'))

    return redirect(url_for('web.ver_producto', id=id))


def dar_de_baja(id):
    producto = ProductoModel.obtener_por_id(id)
    if producto:
        ProductoModel.actualizar(id, {"estado": "inactivo", "activo": False, "updated_at": datetime.utcnow()})
        flash('Producto dado de baja exitosamente', 'success')
    else:
        flash('Producto no encontrado', 'danger')
    return redirect(url_for('web.lista_productos'))


def borrar(id):
    producto = ProductoModel.obtener_por_id(id)
    if not producto:
        flash('Producto no encontrado', 'danger')
        return redirect(url_for('web.lista_productos'))

    if producto.get('vendedor_id') and MARKETPLACE_OK and Vendedor is not None:
        try:
            db = current_app.db
            db[Vendedor.COLLECTION].update_one(
                {'_id': ObjectId(producto['vendedor_id'])},
                {'$inc': {'metricas.productos': -1}}
            )
        except Exception as e:
            print(f"⚠️ No se pudo decrementar métrica: {e}")

    if ProductoModel.eliminar(id):
        flash('Producto eliminado exitosamente', 'success')
    else:
        flash('Producto no encontrado', 'danger')
    return redirect(url_for('web.lista_productos'))


# --- CATÁLOGO Y TIENDA ---

def catalogo():
    """
    Catálogo general público.

    ⭐ Fix: si al filtrar por categoría NO hay productos pero SÍ hay
    productos de marketplace, los muestra igualmente (fallback) para que
    los productos del vendedor sean visibles.
    """
    cat_id = request.args.get('categoria')
    marca_id = request.args.get('marca')
    genero = request.args.get('genero')
    query = request.args.get('q', '')

    categorias_crudas = Categoria.obtener_todas()
    todas = _calcular_jerarquia_categorias(categorias_crudas)

    # ⭐ Cargar TODOS los productos activos (incluye marketplace)
    todos_activos = _obtener_activos_incluyendo_marketplace()

    if cat_id:
        productos = _obtener_activos_por_categoria(cat_id)

        # ⭐ FALLBACK: si el filtro de categoría vacía todo pero hay
        # productos de vendedor, mostrarlos igual (modo "relacionados")
        if not productos:
            productos_mkt = [p for p in todos_activos if p.get('vendedor_id')]
            if productos_mkt:
                print(f"   ℹ️ Categoría '{cat_id}' vacía pero hay "
                      f"{len(productos_mkt)} productos de vendedor → mostrando todos")
                productos = productos_mkt
            else:
                print(f"   ℹ️ Categoría '{cat_id}' vacía y no hay productos de vendedor")
                productos = todos_activos
    else:
        productos = todos_activos

    print(f"📦 catálogo: {len(productos)} productos antes de filtros "
          f"(cat={cat_id}, marca={marca_id}, genero={genero}, q={query!r})")

    # Filtros posteriores
    if marca_id:
        productos = [p for p in productos if str(p.get('marca_id')) == str(marca_id)]
    if genero:
        productos = [
            p for p in productos
            if (p.get('genero') or '').strip().lower() == genero.strip().lower()
        ]
    if query:
        query_lower = query.lower()
        productos = [
            p for p in productos
            if query_lower in (p.get('nombre') or '').lower()
            or query_lower in (p.get('descripcion') or '').lower()
        ]

    # ⭐ Enriquecer con datos del vendedor (badge azul)
    _enriquecer_lista_con_vendedor(productos)

    con_vendedor = sum(1 for p in productos if p.get('vendedor_id'))
    print(f"📦 catálogo FINAL: {len(productos)} productos · {con_vendedor} son de vendedores")

    cat_actual = next((c for c in todas if str(c['_id']) == str(cat_id)), None)
    marcas = Marca.obtener_todas()
    marca_actual = next((m for m in marcas if str(m['_id']) == str(marca_id)), None) if marca_id else None

    ids_visibles = []
    if cat_actual:
        ids_visibles.append(str(cat_actual['_id']))
        if cat_actual.get('padre_id'):
            ids_visibles.append(str(cat_actual['padre_id']))
            padre = next((c for c in todas if str(c['_id']) == str(cat_actual['padre_id'])), None)
            if padre and padre.get('padre_id'):
                ids_visibles.append(str(padre['padre_id']))

    favoritos_ids = _get_favoritos_ids()

    return render_template('tienda/catalogo.html',
                           categorias=todas,
                           productos=productos,
                           categoria_seleccionada=cat_id,
                           cat_actual=cat_actual,
                           ids_visibles=ids_visibles,
                           marca_seleccionada=marca_id,
                           marca_actual=marca_actual,
                           genero_seleccionado=genero,
                           query=query,
                           favoritos_ids=favoritos_ids)


def ver_detalle_producto(id):
    producto = ProductoModel.obtener_por_id(id)
    if not producto:
        return "Producto no encontrado", 404

    if 'variables' in producto and producto['variables']:
        producto['variantes_para_template'] = producto['variables']
    elif 'variantes' in producto and producto['variantes']:
        producto['variantes_para_template'] = producto['variantes']
    else:
        producto['variantes_para_template'] = []

    opiniones = Resena.obtener_por_producto(id)
    producto['opiniones'] = opiniones

    marcas = Marca.obtener_todas()
    marca_actual = next((m for m in marcas if str(m['_id']) == str(producto.get('marca_id'))), None)

    if marca_actual:
        producto['marca_nombre'] = marca_actual.get('nombre_comercial')
        producto['marca'] = marca_actual.get('nombre_comercial')
    else:
        producto['marca_nombre'] = 'Marca no especificada'
        producto['marca'] = 'Marca no especificada'

    categorias_crudas = Categoria.obtener_todas()
    categorias = _calcular_jerarquia_categorias(categorias_crudas)
    categoria_actual = next((c for c in categorias if str(c['_id']) == str(producto.get('categoria_id'))), None)
    producto['categoria'] = categoria_actual.get('nombre') if categoria_actual else producto.get('categoria', '')
    producto['genero'] = producto.get('genero') or (categoria_actual.get('genero') if categoria_actual else '')

    db = current_app.db

    _enriquecer_producto_con_vendedor(producto)

    productos_relacionados = []
    if producto.get('categoria_id'):
        productos_relacionados = [
            p for p in _obtener_activos_incluyendo_marketplace()
            if str(p.get('categoria_id')) == str(producto.get('categoria_id'))
            and str(p['_id']) != str(producto['_id'])
        ][:8]

    ids_relacionados = {str(p['_id']) for p in productos_relacionados}
    productos_complementa = [
        p for p in _obtener_activos_incluyendo_marketplace()
        if str(p['_id']) != str(producto['_id'])
        and str(p['_id']) not in ids_relacionados
    ][:8]

    productos_de_marca = []
    if producto.get('marca_id'):
        productos_de_marca = [
            p for p in _obtener_activos_incluyendo_marketplace()
            if str(p.get('marca_id')) == str(producto.get('marca_id'))
            and str(p['_id']) != str(producto['_id'])
        ][:12]

    productos_mas_vendidos = []
    try:
        pipeline = [
            {"$unwind": "$productos"},
            {"$group": {"_id": "$productos.id", "total_vendido": {"$sum": "$productos.cantidad"}}},
            {"$sort": {"total_vendido": -1}},
            {"$limit": 8}
        ]
        top_ids = [r['_id'] for r in db.ventas.aggregate(pipeline)]
        for pid in top_ids:
            if str(pid) == str(producto['_id']):
                continue
            p = ProductoModel.obtener_por_id(pid)
            if p and p.get('activo', True):
                productos_mas_vendidos.append(p)
    except Exception:
        productos_mas_vendidos = []

    es_favorito = False
    if session.get('user_id'):
        try:
            es_favorito = db.favoritos.find_one({
                'usuario_id': ObjectId(session['user_id']),
                'producto_id': ObjectId(id)
            }) is not None
        except Exception:
            es_favorito = False

    productos_del_vendedor = []
    if producto.get('vendedor_id'):
        try:
            productos_del_vendedor = [
                p for p in _obtener_activos_incluyendo_marketplace()
                if str(p.get('vendedor_id') or '') == str(producto['vendedor_id'])
                and str(p['_id']) != str(producto['_id'])
            ][:8]
            _enriquecer_lista_con_vendedor(productos_del_vendedor)
        except Exception:
            productos_del_vendedor = []

    _enriquecer_lista_con_vendedor(productos_relacionados)
    _enriquecer_lista_con_vendedor(productos_complementa)
    _enriquecer_lista_con_vendedor(productos_de_marca)
    _enriquecer_lista_con_vendedor(productos_mas_vendidos)

    return render_template(
        'tienda/detalle_producto.html',
        producto=producto,
        marcas=marcas,
        productos_relacionados=productos_relacionados,
        productos_complementa=productos_complementa,
        productos_de_marca=productos_de_marca,
        productos_mas_vendidos=productos_mas_vendidos,
        productos_del_vendedor=productos_del_vendedor,
        es_favorito=es_favorito
    )


def enviar_opinion():
    if request.method == 'POST':
        if 'user_id' not in session:
            flash("Debes iniciar sesión para opinar.", "warning")
            return redirect(url_for('web.login'))

        producto_id = request.form.get('producto_id')
        calificacion = request.form.get('calificacion')
        titulo = request.form.get('titulo')
        comentario = request.form.get('comentario')

        if not producto_id or not calificacion or not titulo or not comentario:
            flash("Todos los campos son obligatorios.", "danger")
            return redirect(url_for('web.ver_detalle_producto', id=producto_id))

        archivos = request.files.getlist('fotos')
        lista_archivos = []
        folder = os.path.join(current_app.root_path, 'static', 'uploads', 'resenas')
        os.makedirs(folder, exist_ok=True)

        for archivo in archivos:
            if archivo and archivo.filename:
                nombre_archivo = f"resena_{uuid.uuid4().hex[:8]}_{secure_filename(archivo.filename)}"
                archivo.save(os.path.join(folder, nombre_archivo))
                lista_archivos.append(nombre_archivo)

        compra_verificada = False
        try:
            db = current_app.db
            compras = db.ventas.find_one({
                'usuario_id': session.get('user_id'),
                'productos.id': producto_id
            })
            if compras:
                compra_verificada = True
        except Exception:
            compra_verificada = False

        data = {
            "producto_id": producto_id,
            "usuario_id": session.get('user_id'),
            "usuario_nombre": session.get('nombre', 'Usuario'),
            "calificacion": int(calificacion),
            "titulo": titulo,
            "comentario": comentario,
            "foto_path": lista_archivos,
            "fecha": datetime.now().strftime("%d/%m/%Y"),
            "compra_verificada": compra_verificada,
            "votos_utiles": [],
            "reportes": 0,
            "created_at": datetime.utcnow(),
            "updated_at": datetime.utcnow()
        }

        try:
            Resena.crear(data)
            flash("¡Opinión enviada correctamente!", "success")
        except Exception as e:
            print(f"❌ Error al crear opinión: {e}")
            flash("Error al enviar la opinión.", "danger")

        return redirect(url_for('web.ver_detalle_producto', id=producto_id))

    return redirect(url_for('web.catalogo'))


def listar_por_marca(marca):
    """Lista productos por marca incluyendo marketplace."""
    productos = _obtener_activos_incluyendo_marketplace()

    marcas = Marca.obtener_todas()
    marca_ids_validos = {str(marca)}
    marca_obj = next(
        (m for m in marcas
         if str(m.get('_id')) == str(marca)
         or m.get('nombre_comercial', '').strip().lower() == str(marca).strip().lower()),
        None
    )
    if marca_obj:
        marca_ids_validos.add(str(marca_obj['_id']))

    productos = [
        p for p in productos
        if str(p.get('marca_id') or '') in marca_ids_validos
        or str(p.get('marca') or '').strip().lower() == str(marca).strip().lower()
    ]

    categorias_crudas = Categoria.obtener_todas()
    categorias = _calcular_jerarquia_categorias(categorias_crudas)
    favoritos_ids = _get_favoritos_ids()

    _enriquecer_lista_con_vendedor(productos)

    return render_template('tienda/catalogo.html',
                           productos=productos,
                           categorias=categorias,
                           ids_visibles=[],
                           favoritos_ids=favoritos_ids)


# ================================================================
# ====== FUNCIONES API Y ADMIN ======
# ================================================================

def productos_relacionados(id):
    db = current_app.db
    producto = ProductoModel.obtener_por_id(id)
    if not producto:
        return jsonify({'success': False, 'message': 'Producto no encontrado'}), 404

    categoria = producto.get('categoria_id')
    relacionados = list(db.productos.find({
        '_id': {'$ne': ObjectId(id)},
        'categoria_id': categoria,
        'activo': True
    }).limit(8))

    for p in relacionados:
        p['_id'] = str(p['_id'])
        _enriquecer_producto_con_vendedor(p)

    return jsonify({'success': True, 'productos': relacionados})


def productos_mas_vendidos():
    db = current_app.db
    try:
        pipeline = [
            {"$unwind": "$productos"},
            {"$group": {"_id": "$productos.id", "total_vendido": {"$sum": "$productos.cantidad"}}},
            {"$sort": {"total_vendido": -1}},
            {"$limit": 8}
        ]
        top_ids = [r['_id'] for r in db.ventas.aggregate(pipeline)]
        productos = []
        for pid in top_ids:
            p = ProductoModel.obtener_por_id(pid)
            if p and p.get('activo', True):
                p['_id'] = str(p['_id'])
                _enriquecer_producto_con_vendedor(p)
                productos.append(p)
        return jsonify({'success': True, 'productos': productos})
    except Exception as e:
        return jsonify({'success': True, 'productos': [], 'message': str(e)})


def productos_por_categoria_api(categoria_id):
    db = current_app.db
    productos = list(db.productos.find({
        'categoria_id': ObjectId(categoria_id),
        'activo': True
    }).limit(20))
    for p in productos:
        p['_id'] = str(p['_id'])
        _enriquecer_producto_con_vendedor(p)
    return jsonify({'success': True, 'productos': productos})


def buscar_productos():
    db = current_app.db
    query = request.args.get('q', '')
    if not query:
        return jsonify({'success': True, 'productos': []})
    productos = list(db.productos.find({
        '$or': [
            {'nombre': {'$regex': query, '$options': 'i'}},
            {'descripcion': {'$regex': query, '$options': 'i'}}
        ],
        'activo': True
    }).limit(20))
    for p in productos:
        p['_id'] = str(p['_id'])
        _enriquecer_producto_con_vendedor(p)
    return jsonify({'success': True, 'productos': productos})


def api_productos():
    db = current_app.db
    page = int(request.args.get('page', 1))
    per_page = int(request.args.get('per_page', 20))
    skip = (page - 1) * per_page
    productos = list(db.productos.find({'activo': True}).skip(skip).limit(per_page))
    total = db.productos.count_documents({'activo': True})
    for p in productos:
        p['_id'] = str(p['_id'])
        _enriquecer_producto_con_vendedor(p)
    return jsonify({
        'success': True,
        'productos': productos,
        'total': total,
        'page': page,
        'per_page': per_page,
        'total_pages': (total + per_page - 1) // per_page
    })


def api_producto(id):
    db = current_app.db
    producto = ProductoModel.obtener_por_id(id)
    if not producto:
        return jsonify({'success': False, 'message': 'Producto no encontrado'}), 404
    producto['_id'] = str(producto['_id'])
    _enriquecer_producto_con_vendedor(producto)
    return jsonify({'success': True, 'producto': producto})


def subir_imagen_producto(id):
    if 'user_id' not in session:
        return jsonify({'success': False, 'message': 'Inicia sesión'}), 401
    db = current_app.db
    usuario = db.usuarios.find_one({'_id': ObjectId(session['user_id'])})

    vendedor_actual = _get_vendedor_actual()
    es_admin = usuario and usuario.get('rol') == 'admin'

    if not es_admin and not vendedor_actual:
        return jsonify({'success': False, 'message': 'No autorizado'}), 403

    if 'imagen' not in request.files:
        return jsonify({'success': False, 'message': 'No se envió ninguna imagen'}), 400
    imagen = request.files['imagen']
    if imagen.filename == '':
        return jsonify({'success': False, 'message': 'No se seleccionó ninguna imagen'}), 400

    prod = ProductoModel.obtener_por_id(id)
    if not prod:
        return jsonify({'success': False, 'message': 'Producto no encontrado'}), 404

    if vendedor_actual and not es_admin:
        if str(prod.get('vendedor_id') or '') != str(vendedor_actual['_id']):
            return jsonify({'success': False, 'message': 'Este producto no te pertenece'}), 403

    fotos_actuales = prod.get('fotos', [])
    fotos_nuevas = _save_photos([imagen], prod.get('nombre', 'producto'), fotos_actuales)

    if len(fotos_nuevas) <= len(fotos_actuales):
        return jsonify({'success': False, 'message': 'No se pudo guardar la imagen'}), 400

    db.productos.update_one(
        {'_id': ObjectId(id)},
        {'$set': {'fotos': fotos_nuevas, 'updated_at': datetime.utcnow()}}
    )
    return jsonify({'success': True, 'message': 'Imagen subida correctamente', 'filename': fotos_nuevas[-1]})


def eliminar_imagen_producto(id, index):
    if 'user_id' not in session:
        return jsonify({'success': False, 'message': 'Inicia sesión'}), 401
    db = current_app.db
    usuario = db.usuarios.find_one({'_id': ObjectId(session['user_id'])})

    vendedor_actual = _get_vendedor_actual()
    es_admin = usuario and usuario.get('rol') == 'admin'

    if not es_admin and not vendedor_actual:
        return jsonify({'success': False, 'message': 'No autorizado'}), 403

    producto = ProductoModel.obtener_por_id(id)
    if not producto:
        return jsonify({'success': False, 'message': 'Producto no encontrado'}), 404

    if vendedor_actual and not es_admin:
        if str(producto.get('vendedor_id') or '') != str(vendedor_actual['_id']):
            return jsonify({'success': False, 'message': 'Este producto no te pertenece'}), 403

    fotos = producto.get('fotos', [])
    if index < len(fotos):
        filename = fotos[index]
        filepath = os.path.join(current_app.config['UPLOAD_FOLDER'], 'productos', filename)
        if os.path.exists(filepath):
            os.remove(filepath)
        fotos.pop(index)
        ProductoModel.actualizar(id, {'fotos': fotos, 'updated_at': datetime.utcnow()})
        return jsonify({'success': True, 'message': 'Imagen eliminada correctamente'})
    return jsonify({'success': False, 'message': 'Imagen no encontrada'}), 404


# ================================================================
# ⭐ EXPORTAR PRODUCTOS A EXCEL
# ================================================================

def exportar_productos():
    if not OPENPYXL_OK:
        flash('openpyxl no está instalado. Ejecuta: python -m pip install openpyxl', 'danger')
        return redirect(url_for('web.lista_productos'))

    if 'user_id' not in session:
        flash('Inicia sesión para exportar', 'warning')
        return redirect(url_for('web.login'))

    db = current_app.db
    try:
        usuario = db.usuarios.find_one({'_id': ObjectId(session['user_id'])})
    except Exception:
        flash('Sesión inválida', 'danger')
        return redirect(url_for('web.login'))

    es_admin = usuario and usuario.get('rol') == 'admin'
    vendedor_actual = _get_vendedor_actual()

    if not es_admin and not vendedor_actual:
        flash('No autorizado', 'danger')
        return redirect(url_for('web.dashboard'))

    categorias = {str(c['_id']): c.get('nombre', '') for c in Categoria.obtener_todas()}
    marcas = {str(m['_id']): m.get('nombre_comercial', '') for m in Marca.obtener_todas()}
    temporadas = {str(t['_id']): t.get('nombre', '') for t in Temporada.obtener_todas()}

    filtro = {}
    if vendedor_actual and not es_admin:
        filtro['vendedor_id'] = str(vendedor_actual['_id'])

    productos = list(db.productos.find(filtro))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Productos'

    headers = [
        'nombre', 'descripcion', 'categoria_id', 'categoria_nombre',
        'marca_id', 'marca_nombre', 'temporada_id', 'temporada_nombre',
        'estado', 'activo',
        'vendedor_id', 'vendedor_nombre',
        'sku', 'color', 'talla', 'precio', 'stock',
        'fotos', 'fotos_urls',
        'videos', 'videos_urls',
    ]
    ws.append(headers)

    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color='E10098', end_color='E10098', fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center')

    anchos = {
        'A': 30, 'B': 40, 'C': 22, 'D': 22, 'E': 22, 'F': 22,
        'G': 22, 'H': 22, 'I': 12, 'J': 10,
        'K': 28, 'L': 28,
        'M': 20, 'N': 15, 'O': 12, 'P': 12, 'Q': 10,
        'R': 45, 'S': 60, 'T': 45, 'U': 60,
    }
    for col, ancho in anchos.items():
        ws.column_dimensions[col].width = ancho

    try:
        base_url = request.host_url.rstrip('/')
    except Exception:
        base_url = ''

    for p in productos:
        nombre = p.get('nombre', '')
        descripcion = p.get('descripcion', '')
        categoria_id = str(p.get('categoria_id') or '')
        marca_id = str(p.get('marca_id') or '')
        temporada_id = str(p.get('temporada_id') or '')
        estado = p.get('estado', 'activo')
        activo = p.get('activo', True)
        variables = p.get('variables') or []

        vendedor_id  = str(p.get('vendedor_id') or '')
        vendedor_nom = p.get('vendedor_nombre') or p.get('nombre_tienda') or ''

        fotos_nombres = p.get('fotos') or []
        fotos_str = ' | '.join(fotos_nombres)
        fotos_urls = [
            f"{base_url}/static/uploads/productos/{f}" if base_url else f"/static/uploads/productos/{f}"
            for f in fotos_nombres
        ]
        fotos_urls_str = ' | '.join(fotos_urls)

        videos_nombres = p.get('videos') or []
        videos_str = ' | '.join(videos_nombres)
        videos_urls = [
            f"{base_url}/static/uploads/productos/{f}" if base_url else f"/static/uploads/productos/{f}"
            for f in videos_nombres
        ]
        videos_urls_str = ' | '.join(videos_urls)

        fila_comun = [
            nombre, descripcion,
            categoria_id, categorias.get(categoria_id, ''),
            marca_id, marcas.get(marca_id, ''),
            temporada_id, temporadas.get(temporada_id, ''),
            estado, 'Sí' if activo else 'No',
            vendedor_id, vendedor_nom,
        ]

        if variables:
            for v in variables:
                ws.append(fila_comun + [
                    v.get('sku', ''),
                    v.get('color', ''),
                    v.get('tamano', ''),
                    float(v.get('precio', 0) or 0),
                    int(v.get('stock', 0) or 0),
                    fotos_str, fotos_urls_str,
                    videos_str, videos_urls_str,
                ])
        else:
            ws.append(fila_comun + ['', '', '', 0.0, 0, fotos_str, fotos_urls_str, videos_str, videos_urls_str])

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    fecha = datetime.now().strftime('%Y-%m-%d_%H%M')
    response = make_response(output.getvalue())
    response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    response.headers['Content-Disposition'] = f'attachment; filename=orion-productos-{fecha}.xlsx'
    return response


# ================================================================
# ⭐ IMPORTAR PRODUCTOS DESDE EXCEL
# ================================================================

def importar_productos_excel():
    if not OPENPYXL_OK:
        return jsonify({'success': False, 'message': 'openpyxl no está instalado'}), 500

    if 'user_id' not in session:
        return jsonify({'success': False, 'message': 'Inicia sesión'}), 401

    db = current_app.db
    try:
        usuario = db.usuarios.find_one({'_id': ObjectId(session['user_id'])})
    except Exception:
        return jsonify({'success': False, 'message': 'Sesión inválida'}), 401

    es_admin = usuario and usuario.get('rol') == 'admin'
    vendedor_actual = _get_vendedor_actual()

    if not es_admin and not vendedor_actual:
        return jsonify({'success': False, 'message': 'No autorizado'}), 403

    if 'archivo_excel' not in request.files:
        return jsonify({'success': False, 'message': 'No se envió ningún archivo'}), 400

    archivo = request.files['archivo_excel']
    if not archivo or archivo.filename == '':
        return jsonify({'success': False, 'message': 'No se seleccionó ningún archivo'}), 400

    ext = os.path.splitext(archivo.filename)[1].lower()
    if ext not in ['.xlsx', '.xls']:
        return jsonify({'success': False, 'message': 'Solo .xlsx'}), 400

    actualizar_existentes = request.form.get('actualizar_existentes') in ['1', 'true', 'on', 'yes']

    try:
        contenido = archivo.read()
        wb = load_workbook(filename=BytesIO(contenido), data_only=True)
        sheet = wb.active
    except Exception as e:
        return jsonify({'success': False, 'message': f'Error: {str(e)}'}), 400

    headers = []
    for row in sheet.iter_rows(min_row=1, max_row=1, values_only=True):
        headers = [str(h).strip().lower() if h else '' for h in row]
        break

    if not headers:
        return jsonify({'success': False, 'message': 'Archivo vacío'}), 400

    mapa_columnas = {}
    alias = {
        'nombre':        ['nombre', 'name', 'producto'],
        'categoria_id':  ['categoria_id', 'categoria', 'category_id', 'category'],
        'marca_id':      ['marca_id', 'marca', 'brand_id', 'brand'],
        'temporada_id':  ['temporada_id', 'temporada', 'season_id', 'season'],
        'sku':           ['sku', 'codigo', 'code'],
        'color':         ['color'],
        'talla':         ['talla', 'tamano', 'size'],
        'precio':        ['precio', 'price', 'precio_unitario'],
        'stock':         ['stock', 'cantidad', 'quantity', 'existencia'],
        'descripcion':   ['descripcion', 'description', 'detalle'],
        'estado':        ['estado', 'status', 'activo'],
        'fotos':         ['fotos', 'fotos_nombres', 'imagenes', 'imagen'],
        'fotos_urls':    ['fotos_urls', 'urls_imagenes', 'imagenes_urls', 'urls', 'imagenes_url'],
        'videos':        ['videos', 'videos_nombres', 'videos_locales'],
        'videos_urls':   ['videos_urls', 'urls_videos', 'videos_url'],
    }

    for campo, posibles in alias.items():
        for i, h in enumerate(headers):
            if h in posibles:
                mapa_columnas[campo] = i
                break

    if 'nombre' not in mapa_columnas:
        return jsonify({'success': False, 'message': 'Falta columna "nombre"'}), 400

    def _separar_lista(valor):
        if valor is None:
            return []
        s = str(valor).strip()
        if not s:
            return []
        for sep in [',', ';', '\n']:
            s = s.replace(sep, '|')
        return [x.strip() for x in s.split('|') if x.strip()]

    def _descargar_archivo(url, nombre_producto, index, tipo='img'):
        try:
            upload_path = _get_upload_folder()
            ext = os.path.splitext(url.split('?')[0])[1].lower()
            if tipo == 'img':
                if ext not in EXTENSIONES_IMAGEN:
                    ext = '.jpg'
                prefijo_tipo = 'img'
            else:
                if ext not in EXTENSIONES_VIDEO:
                    ext = '.mp4'
                prefijo_tipo = 'vid'

            prefijo = "".join(c if c.isalnum() else "_" for c in nombre_producto.lower())[:25]
            nombre_archivo = f"{prefijo}_{prefijo_tipo}_imp_{index}_{os.urandom(2).hex()}{ext}"

            req = urllib.request.Request(url, headers={
                'User-Agent': 'Mozilla/5.0 (compatible; OrionBot/1.0)'
            })
            with urllib.request.urlopen(req, timeout=30) as resp:
                contenido = resp.read()
                max_size = 100 * 1024 * 1024 if tipo == 'vid' else 15 * 1024 * 1024
                if len(contenido) > max_size:
                    return None
                with open(os.path.join(upload_path, nombre_archivo), 'wb') as f:
                    f.write(contenido)
            return nombre_archivo
        except Exception as e:
            print(f"⚠️ Error descargando {tipo} {url}: {e}")
            return None

    def _validar_local(nombre):
        upload_path = _get_upload_folder()
        return os.path.exists(os.path.join(upload_path, nombre))

    productos_agrupados = {}
    total_filas = 0
    filas_vacias = 0
    errores = []

    for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not row or all(c is None or str(c).strip() == '' for c in row):
            filas_vacias += 1
            continue

        total_filas += 1

        def get_col(campo, default=''):
            idx = mapa_columnas.get(campo)
            if idx is None or idx >= len(row):
                return default
            v = row[idx]
            return v if v is not None else default

        nombre = str(get_col('nombre', '')).strip()
        if not nombre:
            errores.append(f'Fila {row_idx}: falta nombre')
            continue

        sku = str(get_col('sku', '')).strip()
        color = str(get_col('color', '')).strip()
        talla = str(get_col('talla', '')).strip()
        descripcion = str(get_col('descripcion', '')).strip()
        categoria_id = str(get_col('categoria_id', '')).strip()
        marca_id = str(get_col('marca_id', '')).strip()
        temporada_id = str(get_col('temporada_id', '')).strip()
        estado = str(get_col('estado', '')).strip() or 'activo'

        try:
            precio = float(get_col('precio', 0) or 0)
        except (ValueError, TypeError):
            precio = 0.0

        try:
            stock = int(float(get_col('stock', 0) or 0))
        except (ValueError, TypeError):
            stock = 0

        fotos_nombres = _separar_lista(get_col('fotos', ''))
        fotos_urls = _separar_lista(get_col('fotos_urls', ''))
        videos_nombres = _separar_lista(get_col('videos', ''))
        videos_urls = _separar_lista(get_col('videos_urls', ''))

        categoria_id = categoria_id or None
        marca_id = marca_id or None
        temporada_id = temporada_id or None

        if nombre not in productos_agrupados:
            productos_agrupados[nombre] = {
                'nombre': nombre,
                'descripcion': descripcion,
                'categoria_id': categoria_id,
                'marca_id': marca_id,
                'temporada_id': temporada_id,
                'estado': estado,
                'variantes': [],
                'fotos_nombres': fotos_nombres,
                'fotos_urls': fotos_urls,
                'videos_nombres': videos_nombres,
                'videos_urls': videos_urls,
            }

        productos_agrupados[nombre]['variantes'].append({
            'sku': sku or f"{nombre[:6].upper()}-{len(productos_agrupados[nombre]['variantes']) + 1}",
            'color': color,
            'tamano': talla,
            'precio': precio,
            'stock': stock,
        })

    creados = 0
    actualizados = 0
    omitidos = 0
    total_fotos_descargadas = 0
    total_videos_descargados = 0

    for nombre_prod, data_prod in productos_agrupados.items():
        try:
            existente = db.productos.find_one({'nombre': nombre_prod})

            if vendedor_actual and not es_admin and existente:
                if str(existente.get('vendedor_id') or '') != str(vendedor_actual['_id']):
                    errores.append(f'"{nombre_prod}": no es tu producto')
                    omitidos += 1
                    continue

            fotos_finales = []
            for fn in data_prod.get('fotos_nombres', []):
                if _validar_local(fn):
                    fotos_finales.append(fn)
                else:
                    errores.append(f'"{nombre_prod}": foto "{fn}" no existe')
            for i, url in enumerate(data_prod.get('fotos_urls', [])):
                if not url.startswith('http') or len(fotos_finales) >= MAX_FOTOS:
                    continue
                nombre_guardado = _descargar_archivo(url, nombre_prod, i + 1, 'img')
                if nombre_guardado:
                    fotos_finales.append(nombre_guardado)
                    total_fotos_descargadas += 1
            fotos_finales = fotos_finales[:MAX_FOTOS]

            videos_finales = []
            for vn in data_prod.get('videos_nombres', []):
                if _validar_local(vn):
                    videos_finales.append(vn)
                else:
                    errores.append(f'"{nombre_prod}": video "{vn}" no existe')
            for i, url in enumerate(data_prod.get('videos_urls', [])):
                if not url.startswith('http') or len(videos_finales) >= MAX_VIDEOS:
                    continue
                nombre_guardado = _descargar_archivo(url, nombre_prod, i + 1, 'vid')
                if nombre_guardado:
                    videos_finales.append(nombre_guardado)
                    total_videos_descargados += 1
            videos_finales = videos_finales[:MAX_VIDEOS]

            if existente and actualizar_existentes:
                update_data = {
                    'descripcion': data_prod['descripcion'] or existente.get('descripcion', ''),
                    'categoria_id': data_prod['categoria_id'] or existente.get('categoria_id'),
                    'marca_id': data_prod['marca_id'] or existente.get('marca_id'),
                    'temporada_id': data_prod['temporada_id'] or existente.get('temporada_id'),
                    'estado': data_prod['estado'] or existente.get('estado', 'activo'),
                    'variables': data_prod['variantes'],
                    'updated_at': datetime.utcnow(),
                }
                if fotos_finales:
                    update_data['fotos'] = fotos_finales
                if videos_finales:
                    update_data['videos'] = videos_finales

                if vendedor_actual:
                    update_data['vendedor_id']     = str(vendedor_actual['_id'])
                    update_data['vendedor_nombre'] = vendedor_actual.get('nombre_tienda', '')
                    update_data['nombre_tienda']   = vendedor_actual.get('nombre_tienda', '')
                    update_data['tienda_nombre']   = vendedor_actual.get('nombre_tienda', '')

                db.productos.update_one({'_id': existente['_id']}, {'$set': update_data})
                actualizados += 1

            elif existente and not actualizar_existentes:
                omitidos += 1

            else:
                nuevo = {
                    'nombre': data_prod['nombre'],
                    'descripcion': data_prod['descripcion'],
                    'categoria_id': data_prod['categoria_id'],
                    'marca_id': data_prod['marca_id'],
                    'temporada_id': data_prod['temporada_id'],
                    'estado': data_prod['estado'],
                    'activo': True,
                    'variables': data_prod['variantes'],
                    'fotos': fotos_finales,
                    'videos': videos_finales,
                    'created_at': datetime.utcnow(),
                    'updated_at': datetime.utcnow(),
                }

                if vendedor_actual:
                    nuevo['vendedor_id']     = str(vendedor_actual['_id'])
                    nuevo['vendedor_nombre'] = vendedor_actual.get('nombre_tienda', '')
                    nuevo['nombre_tienda']   = vendedor_actual.get('nombre_tienda', '')
                    nuevo['tienda_nombre']   = vendedor_actual.get('nombre_tienda', '')

                db.productos.insert_one(nuevo)
                creados += 1

        except Exception as e:
            errores.append(f'"{nombre_prod}": {str(e)}')

    partes = []
    if creados: partes.append(f'{creados} creado{"s" if creados != 1 else ""}')
    if actualizados: partes.append(f'{actualizados} actualizado{"s" if actualizados != 1 else ""}')
    if omitidos: partes.append(f'{omitidos} omitido{"s" if omitidos != 1 else ""}')
    if total_fotos_descargadas: partes.append(f'{total_fotos_descargadas} img descargada{"s" if total_fotos_descargadas != 1 else ""}')
    if total_videos_descargados: partes.append(f'{total_videos_descargados} video descargado{"s" if total_videos_descargados != 1 else ""}')

    mensaje = ', '.join(partes) if partes else 'No se procesaron productos'

    return jsonify({
        'success': True,
        'message': mensaje,
        'detalle': {
            'creados': creados,
            'actualizados': actualizados,
            'omitidos': omitidos,
            'total_filas': total_filas,
            'filas_vacias': filas_vacias,
            'productos_agrupados': len(productos_agrupados),
            'imagenes_descargadas': total_fotos_descargadas,
            'videos_descargados': total_videos_descargados,
            'errores': errores[:10],
        }
    })


# ================================================================
# ⭐ AUTOCOMPLETADO DE BÚSQUEDA
# ================================================================
def autocompletar_productos():
    try:
        q = (request.args.get('q') or '').strip()
        if not q or len(q) < 2:
            return jsonify({"success": True, "query": q, "total": 0, "resultados": []})

        db = current_app.db
        limite = 8
        patron = {'$regex': q, '$options': 'i'}

        cursor = db.productos.find({
            'activo': {'$ne': False},
            '$or': [
                {'nombre': patron},
                {'descripcion': patron}
            ]
        }).limit(limite)

        resultados = []
        for p in cursor:
            precio = 0.0
            variables = p.get('variables') or []
            variantes_con_stock = [v for v in variables if (v.get('stock') or 0) > 0]
            if variantes_con_stock:
                precio = float(variantes_con_stock[0].get('precio') or 0)
            elif variables:
                precio = float(variables[0].get('precio') or 0)
            else:
                precio = float(p.get('precio') or 0)

            foto_url = ''
            fotos = p.get('fotos') or []
            if fotos:
                foto_url = url_for('static', filename='uploads/productos/' + fotos[0])

            marca_nombre = ''
            marca_id = p.get('marca_id')
            if marca_id:
                try:
                    m = db.marcas.find_one({'_id': ObjectId(str(marca_id))})
                    if m:
                        marca_nombre = m.get('nombre_comercial') or m.get('nombre') or ''
                except Exception:
                    marca_nombre = p.get('marca_nombre') or ''

            categoria_nombre = ''
            cat_id = p.get('categoria_id')
            if cat_id:
                try:
                    c = db.categorias.find_one({'_id': ObjectId(str(cat_id))})
                    if c:
                        categoria_nombre = c.get('nombre') or ''
                except Exception:
                    categoria_nombre = p.get('categoria') or ''

            _enriquecer_producto_con_vendedor(p)

            nombre_vendedor = _nombre_vendedor_display(p)
            badge_vendedor  = _truncar_nombre_vendedor(nombre_vendedor, BADGE_VENDEDOR_MAX_LEN)

            resultados.append({
                'id': str(p['_id']),
                'nombre': p.get('nombre') or '',
                'precio': round(precio, 2),
                'marca': marca_nombre,
                'categoria': categoria_nombre,
                'vendedor': nombre_vendedor,
                'vendedor_badge': badge_vendedor,
                'es_marketplace': bool(p.get('vendedor_id')),
                'foto': foto_url,
                'url': url_for('web.ver_detalle_producto', id=str(p['_id']))
            })

        return jsonify({
            "success": True,
            "query": q,
            "total": len(resultados),
            "resultados": resultados
        })

    except Exception as e:
        print(f"❌ Error en autocompletar_productos: {e}")
        import traceback
        traceback.print_exc()
        return jsonify({"success": False, "error": str(e), "resultados": []}), 500


# ================================================================
# ⭐ DEBUG: diagnóstico
# ================================================================
def debug_productos():
    """⚠️ Quitar en producción. Visitar /debug/productos"""
    db = current_app.db

    total           = db.productos.count_documents({})
    activos         = db.productos.count_documents({'activo': True})
    inactivos       = db.productos.count_documents({'activo': False})
    con_vendedor    = db.productos.count_documents({'vendedor_id': {'$exists': True, '$ne': None}})
    sin_vendedor    = db.productos.count_documents({
        '$or': [
            {'vendedor_id': {'$exists': False}},
            {'vendedor_id': None},
        ]
    })

    estados_agg = list(db.productos.aggregate([
        {'$group': {'_id': '$estado', 'count': {'$sum': 1}}},
        {'$sort': {'count': -1}},
    ]))

    muestra = list(db.productos.find(
        {'vendedor_id': {'$exists': True, '$ne': None}},
        {'nombre': 1, 'vendedor_id': 1, 'vendedor_nombre': 1,
         'estado': 1, 'activo': 1, 'categoria_id': 1}
    ).limit(5))

    for p in muestra:
        p['_id'] = str(p['_id'])
        p['categoria_id'] = str(p.get('categoria_id')) if p.get('categoria_id') else None

    try:
        del_modelo = ProductoModel.obtener_activos()
        modelo_total     = len(del_modelo)
        modelo_con_vend  = sum(1 for p in del_modelo if p.get('vendedor_id'))
    except Exception as e:
        modelo_total = f"ERROR: {e}"
        modelo_con_vend = "ERROR"

    try:
        del_helper = _obtener_activos_incluyendo_marketplace()
        helper_total     = len(del_helper)
        helper_con_vend  = sum(1 for p in del_helper if p.get('vendedor_id'))
    except Exception as e:
        helper_total = f"ERROR: {e}"
        helper_con_vend = "ERROR"

    return jsonify({
        'resumen': {
            'total_en_db':           total,
            'activos_True':          activos,
            'activos_False':         inactivos,
            'activos_ausentes':      total - activos - inactivos,
            'con_vendedor_id':       con_vendedor,
            'sin_vendedor_id':       sin_vendedor,
        },
        'estados_usados': [
            {'estado': e['_id'], 'count': e['count']} for e in estados_agg
        ],
        'modelo_obtener_activos': {
            'total':            modelo_total,
            'con_vendedor_id':  modelo_con_vend,
        },
        'helper_robusto': {
            'total':            helper_total,
            'con_vendedor_id':  helper_con_vend,
        },
        'muestra_vendedores': muestra,
    })


def debug_categorias():
    """⚠️ Quitar en producción. Visitar /debug/categorias"""
    db = current_app.db

    categorias = list(db.categorias.find({}, {'nombre': 1, 'padre_id': 1}))
    categorias_info = [
        {'_id': str(c['_id']), 'nombre': c.get('nombre'), 'padre_id': str(c.get('padre_id') or '')}
        for c in categorias
    ]

    agg = list(db.productos.aggregate([
        {'$group': {'_id': '$categoria_id', 'count': {'$sum': 1}}},
        {'$sort': {'count': -1}},
        {'$limit': 30},
    ]))

    def _norm(x):
        try:
            return str(x) if x is not None else None
        except Exception:
            return repr(x)

    dist = [{'categoria_id': _norm(r['_id']), 'count': r['count']} for r in agg]

    sin_cat = db.productos.count_documents({'categoria_id': {'$exists': False}})
    con_cat = db.productos.count_documents({'categoria_id': {'$exists': True, '$ne': None}})

    tipos = {}
    for p in db.productos.find({'categoria_id': {'$exists': True}}).limit(200):
        t = type(p.get('categoria_id')).__name__
        tipos[t] = tipos.get(t, 0) + 1

    muestra = list(db.productos.find(
        {},
        {'nombre': 1, 'categoria_id': 1, 'categoria': 1, 'vendedor_id': 1}
    ).limit(10))
    for p in muestra:
        p['_id'] = str(p['_id'])
        p['categoria_id'] = _norm(p.get('categoria_id'))

    return jsonify({
        'categorias_en_bd': categorias_info,
        'total_categorias': len(categorias),
        'productos_con_categoria_id': con_cat,
        'productos_sin_categoria_id': sin_cat,
        'tipos_de_dato_categoria_id': tipos,
        'distribucion_categoria_id_en_productos': dist,
        'muestra_productos': muestra,
    })


def debug_catalogo():
    """
    ⚠️ Diagnóstico. Visitar /debug/catalogo
    Muestra TODO lo que hay en productos: id, nombre, vendedor, categoría, activo, estado.
    """
    db = current_app.db

    productos = list(db.productos.find({}))
    salida = []
    for p in productos:
        salida.append({
            '_id':               str(p.get('_id')),
            'nombre':            p.get('nombre'),
            'vendedor_id':       str(p.get('vendedor_id')) if p.get('vendedor_id') else None,
            'vendedor_nombre':   p.get('vendedor_nombre'),
            'categoria_id':      str(p.get('categoria_id')) if p.get('categoria_id') else None,
            'categoria_id_tipo': type(p.get('categoria_id')).__name__ if p.get('categoria_id') else None,
            'categoria':         p.get('categoria'),
            'marca_id':          str(p.get('marca_id')) if p.get('marca_id') else None,
            'activo':            p.get('activo'),
            'estado':            p.get('estado'),
            'es_marketplace':    bool(p.get('vendedor_id')),
        })

    cats = list(db.categorias.find({}))
    categorias = [{
        '_id':       str(c['_id']),
        'nombre':    c.get('nombre'),
        'padre_id':  str(c.get('padre_id')) if c.get('padre_id') else None,
    } for c in cats]

    return jsonify({
        'total_productos': len(salida),
        'productos': salida,
        'total_categorias': len(categorias),
        'categorias': categorias,
    })


# ================================================================
# 💳 MSI — MESES SIN INTERESES
# ================================================================

def _calcular_msi(precio):
    from app.config.msi_config import MSI_OPCIONES
    precio = float(precio or 0)
    opciones = []

    for op in MSI_OPCIONES:
        if precio < op["min"]:
            continue

        meses = op["meses"]
        tasa = op["tasa_interes"]

        if tasa == 0:
            mensualidad = precio / meses
            total = precio
        else:
            tasa_mensual = (tasa / 100) / 12
            if tasa_mensual == 0:
                mensualidad = precio / meses
            else:
                factor = (tasa_mensual * (1 + tasa_mensual) ** meses) / ((1 + tasa_mensual) ** meses - 1)
                mensualidad = precio * factor
            total = mensualidad * meses

        opciones.append({
            "meses": meses,
            "etiqueta": op["etiqueta"],
            "min": op["min"],
            "tipo": op["tipo"],
            "bancos": op["bancos"],
            "mensualidad": round(mensualidad, 2),
            "total": round(total, 2),
            "tasa": tasa,
        })

    return opciones


def _mejor_msi(precio):
    opciones = _calcular_msi(precio)
    sin_interes = [o for o in opciones if o["tipo"] == "sin_interes"]
    if sin_interes:
        return max(sin_interes, key=lambda o: o["meses"])
    return None


def api_calcular_msi():
    from app.config.msi_config import TARJETAS_TIENDA, MSI_LEGAL
    try:
        precio_raw = request.args.get("precio", "0")
        try:
            precio = float(precio_raw)
        except (ValueError, TypeError):
            precio = 0.0

        opciones = _calcular_msi(precio)
        mejor = _mejor_msi(precio)

        return jsonify({
            "success": True,
            "precio": round(precio, 2),
            "opciones": opciones,
            "mejor_opcion": mejor,
            "tarjetas_tienda": TARJETAS_TIENDA,
            "legal": MSI_LEGAL,
        })
    except Exception as e:
        print(f"❌ Error en api_calcular_msi: {e}")
        import traceback
        traceback.print_exc()
        return jsonify({"success": False, "error": str(e), "opciones": []}), 500